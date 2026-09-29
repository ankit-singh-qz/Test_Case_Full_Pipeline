"""Write every generated test case to an Excel workbook.

Ready cases and blocked cases are separate sheets. Each case is one row;
its steps stay in the Steps, Test data, and Expected result cells.
A Coverage sheet is added when a summary is supplied.

Does not call an LLM. Reads the same test_case_sections dict main.py
already writes to test_case_state.json.
"""
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from orchestrator import SECTION_ORDER, SECTION_TITLES

HEADER_FILL = PatternFill("solid", fgColor="122A54")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")

CASE_HEADERS = [
    "Category",
    "Category name",
    "Test ID",
    "Scenario ID",
    "Title",
    "Objective",
    "Status",
    "Priority",
    "Trace to",
    "Automatable",
    "Blocked By Issue",
    "Reason",
    "Preconditions",
    "Steps",
    "Test data",
    "Expected result",
    "Overall Expected Result",
]

CASE_WIDTHS = [12, 28, 16, 14, 42, 48, 12, 10, 36, 14, 20, 48, 48, 48, 36, 48, 48]


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, list):
        return "\n".join(str(item) for item in value)
    return str(value)


def _category_code(technique: str) -> str:
    """'Positive Testing (PT)' -> 'PT'."""
    if technique.endswith(")") and "(" in technique:
        return technique[technique.rfind("(") + 1:-1]
    return technique


def _format_step_field(steps, key: str) -> str:
    """One numbered line per step, so Steps, Test data, and Expected result line up."""
    if not steps:
        return ""
    lines = []
    for step in steps:
        if not isinstance(step, dict):
            continue
        number = step.get("step_no", "?")
        value = step.get(key)
        if key == "action":
            lines.append(f"{number}. {value or ''}")
        elif value not in (None, ""):
            lines.append(f"{number}. {_text(value)}")
        else:
            lines.append(f"{number}.")
    return "\n".join(lines)


def load_trace_refs(test_design_path: str) -> dict:
    """scenario_id -> source refs (FR-01, fdd_acceptance_scenarios[0], ...).

    Those refs live on the test-design scenario, not on the test case.
    """
    if not test_design_path or not os.path.exists(test_design_path):
        return {}
    with open(test_design_path, encoding="utf-8") as f:
        data = json.load(f)
    traces = {}
    for section in (data.get("test_sections") or {}).values():
        if not isinstance(section, dict):
            continue
        for scenario in section.get("scenarios") or []:
            if not isinstance(scenario, dict) or not scenario.get("scenario_id"):
                continue
            traces[scenario["scenario_id"]] = _text(scenario.get("src_refs"))
    return traces


def _style_sheet(ws, widths) -> None:
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_test_cases_excel(
    path: str,
    test_case_sections: dict,
    coverage_summary: dict = None,
    trace_by_scenario: dict = None,
) -> None:
    wb = Workbook()

    ready_sheet = wb.active
    ready_sheet.title = "Ready"
    ready_sheet.append(CASE_HEADERS)

    blocked_sheet = wb.create_sheet("Blocked")
    blocked_sheet.append(CASE_HEADERS)

    for section_id in SECTION_ORDER:
        section = (test_case_sections or {}).get(section_id) or {}
        technique = SECTION_TITLES[section_id]
        category = _category_code(technique)
        for case in section.get("test_cases") or []:
            if not isinstance(case, dict):
                continue
            case_steps = case.get("steps") or []
            scenario_id = case.get("scenario_id") or ""
            trace_to = (trace_by_scenario or {}).get(scenario_id, "")
            row = [
                category,
                technique,
                case.get("test_case_id") or "",
                scenario_id,
                case.get("title") or "",
                case.get("objective") or "",
                case.get("status") or "",
                case.get("priority") or "",
                trace_to,
                _text(case.get("automatable")),
                _text(case.get("blocked_by_issue")),
                case.get("reason") or "",
                _text(case.get("preconditions")),
                _format_step_field(case_steps, "action"),
                _format_step_field(case_steps, "test_data"),
                _format_step_field(case_steps, "expected_result"),
                case.get("overall_expected_result") or "",
            ]
            if str(case.get("status") or "").lower() == "blocked":
                blocked_sheet.append(row)
            else:
                ready_sheet.append(row)

    _style_sheet(ready_sheet, CASE_WIDTHS)
    _style_sheet(blocked_sheet, CASE_WIDTHS)

    if coverage_summary:
        summary = wb.create_sheet("Coverage", 0)
        summary.append(["Metric", "Value"])
        summary.append(["Total", coverage_summary.get("total", 0)])
        summary.append(["Ready", coverage_summary.get("ready", 0)])
        summary.append(["Blocked", coverage_summary.get("blocked", 0)])
        summary.append(["Ready and automatable", coverage_summary.get("ready_automatable", 0)])
        by_priority = coverage_summary.get("ready_by_priority") or {}
        summary.append([])
        summary.append(["Ready by priority", "Count"])
        for priority in ("P1", "P2", "P3"):
            summary.append([priority, by_priority.get(priority, 0)])
        summary.append([])
        summary.append(["Technique", "Total", "Ready", "Blocked"])
        by_technique = coverage_summary.get("by_technique") or {}
        for section_id in SECTION_ORDER:
            row = by_technique.get(section_id) or {}
            summary.append([
                row.get("title") or SECTION_TITLES[section_id],
                row.get("total", 0),
                row.get("ready", 0),
                row.get("blocked", 0),
            ])
        summary.append([])
        summary.append(["Blocked by issue", "Test cases"])
        for issue_id, count in sorted((coverage_summary.get("by_blocked_issue") or {}).items()):
            summary.append([issue_id, count])
        _style_sheet(summary, [36, 18, 14, 14])

    wb.save(path)


def xlsx_beside_state(state_path: str) -> str:
    """Workbook path in the same folder as the state JSON."""
    return os.path.join(os.path.dirname(os.path.abspath(state_path)), "generated_test_cases.xlsx")


def _export_state_file(state_path: str, xlsx_path: str) -> None:
    with open(state_path, encoding="utf-8") as f:
        data = json.load(f)
    write_test_cases_excel(
        xlsx_path,
        data.get("test_case_sections") or {},
        data.get("coverage_summary"),
        load_trace_refs(data.get("source_test_design_path") or ""),
    )


def export_output_folders(output_dir: str = "output") -> list[str]:
    """Write one workbook into each run folder, plus the stable output folder.

    Each workbook is built from the test_case_state.json that already lives
    in that folder, so older runs get the cases from that run.
    """
    written = []
    locked = []
    if not os.path.isdir(output_dir):
        return written

    jobs = []
    stable_state = os.path.join(output_dir, "test_case_state.json")
    if os.path.exists(stable_state):
        jobs.append(stable_state)
    for name in sorted(os.listdir(output_dir)):
        run_dir = os.path.join(output_dir, name)
        state_path = os.path.join(run_dir, "test_case_state.json")
        if name.startswith("run_") and os.path.isdir(run_dir) and os.path.exists(state_path):
            jobs.append(state_path)

    for state_path in jobs:
        xlsx_path = xlsx_beside_state(state_path)
        try:
            _export_state_file(state_path, xlsx_path)
        except PermissionError:
            locked.append(xlsx_path)
            continue
        written.append(xlsx_path)
    for path in locked:
        print(f"Could not overwrite (close the file in Excel): {path}")
    return written


if __name__ == "__main__":
    if len(sys.argv) > 2:
        state_path, xlsx_path = sys.argv[1], sys.argv[2]
        if not os.path.exists(state_path):
            print(f"Could not find test case state at: {state_path}")
            sys.exit(1)
        _export_state_file(state_path, xlsx_path)
        print(f"Excel written to {xlsx_path}")
    elif len(sys.argv) > 1:
        state_path = sys.argv[1]
        if not os.path.exists(state_path):
            print(f"Could not find test case state at: {state_path}")
            sys.exit(1)
        xlsx_path = xlsx_beside_state(state_path)
        _export_state_file(state_path, xlsx_path)
        print(f"Excel written to {xlsx_path}")
    else:
        written = export_output_folders("output")
        if not written:
            print("No test_case_state.json found under output/")
            sys.exit(1)
        for path in written:
            print(f"Excel written to {path}")
